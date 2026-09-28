"""Dependency smoke: every document format Open WebUI reads, uploaded through the API.

An upload with `process=true` runs through `retrieval/loaders/main.py`, which hands each format to a
third-party library: .pdf to pypdf, .docx to docx2txt, .pptx, .xlsx, .xls, .xml, .rst, .epub and
.odt to unstructured's partitioners (python-pptx; pandas on openpyxl or xlrd behind a msoffcrypto
encryption check, or on pyxlsb for a binary .xlsb workbook sent as an Excel file; pypandoc and the
pandoc binary), .html to BeautifulSoup, plain text through chardet's encoding hint, and every result
through ftfy, which repairs mojibake and drops control characters. With `PDF_EXTRACT_IMAGES` on, a
PDF's images are opened by Pillow (or, kept as raw pixels Pillow cannot open, turned into a picture
by pypdf first) and read by rapidocr on onnxruntime and OpenCV, built from the models its wheel
carries (that instance has every download refused); an image without text adds nothing, so a blank
scan is reported as empty. A dependency bump that breaks one of those paths fails the upload or
loses the text, which `GET /api/v1/files/{id}/data/content` shows. With the Azure Document
Intelligence engine a PDF goes to azure-ai-documentintelligence instead, here against a local
stand-in of the analyze API: the key header, the markdown output format and the polled result are
what it relies on (twin of unit/deps/test_azure_ai_documentintelligence.py). The library contracts
are in unit/deps/.

A workbook is read sheet by sheet in order, numbers included: unstructured's `partition_xlsx`
hands it to pandas' `read_excel` on openpyxl. `pip install open-webui` leaves `unstructured` out,
and an instance booted that way (`harness.missing_packages`) reads spreadsheets with pandas
itself (`ExcelFile`, `read_excel` per sheet on openpyxl or xlrd, `to_string` without the row
index) and slides with python-pptx (twin of unit/deps/test_pandas.py and
unit/deps/test_openpyxl.py). A legacy .xls reaches pandas on xlrd either way: every sheet, its
numbers and its dates, and on the pandas path xlrd's own reason for a workbook it cannot parse
(twin of unit/deps/test_xlrd.py).

Legacy Word (.doc) and PowerPoint (.ppt) files go to unstructured's `partition_doc` and
`partition_ppt`, which convert them with LibreOffice, and a modern file under a legacy name is
told apart by `detect_filetype` and read without it. An Outlook message goes to `partition_msg`,
which reads its body (twin of unit/deps/test_unstructured.py); `harness/outlook_message.py` builds
one.

A PDF is read page by page, each page keeping its label from the page label tree and the
document's title, author and creation date from its info dictionary, or as one document split
at page breaks in single mode. pypdf opens a PDF locked only against editing and refuses one
that needs a password to open, and anything else named .pdf fails its upload. Without the
pandoc program, a document unstructured converts through pypandoc names the missing program.
pandas keeps a workbook's text columns in pyarrow arrays, with unstructured and without it.

Windows-1251 Cyrillic is decoded with the codec chardet names, which ftfy could not repair
after a latin-1 fallback.

EUC-KR and Shift-JIS text were decoded as GB18030 mojibake (issue #31352, fix 3c47f0d7e, PR
#31356): chardet 7.4.3 says CP949 for EUC-KR, which the codec map in `_detect_text_encoding`
lacked, and Shift-JIS was missing from its try order. Both cases pass on dev efe63bd34 and fail
with 3c47f0d7e reverted.

Discriminates: passes on dev bbfa876af (.rst, .epub and .odt with a pandoc binary on PATH). One
backend copy broke pypdf's `extract_text`, `docx2txt.process`, the xlsx, rst and epub partitions and
`chardet.detect`; another broke rapidocr's `RapidOCR`, the pptx, xml and odt partitions,
BeautifulSoup's `get_text` and `ftfy.fix_text`. Each copy turned exactly its own formats red and
left the others green. Mapping cp949 in a third copy makes the EUC-KR case pass. A fourth copy whose
`chardet.detect` names no encoding fails the Big5, EUC-KR, Shift-JIS and Windows-1251 cases. On dev
ef67cc3fa, dropping `output_content_format='markdown'` from the Document Intelligence loader fails
its test; msoffcrypto's `OfficeFile` answering "not encrypted" (patched in at import) fails the
password-protected workbook test, one that accepts any bytes fails the test of a file that only
claims to be a workbook, and OpenCV's `minAreaRect` answering an empty box or onnxruntime refusing
to build a session fails the PDF image test; `RapidOCR` answering every image with the same text
fails both OCR tests, as does one that fetches its models before building. pyxlsb reading every
number as zero fails the binary workbook numbers test, and its strings read as empty fail the .xlsb
upload. A docx2txt result cut to its first paragraph or to ASCII fails the Word paragraphs test,
`ftfy.fix_text` left out fails the mojibake, smart quote and control character tests, `fix_text`
without `unescape_html=False` fails the literal entity test and one that fails on empty text fails
the blank page test (and the empty Word document one). An openpyxl whose `load_workbook` raises
fails both workbook uploads; the pandas loader printing the row index or reading only the first
sheet fails the pandas test, a python-pptx loader that skips text frames fails the slides test, and
a PDF loader that no longer hands an image Pillow cannot open (`UnidentifiedImageError`) to pypdf
fails the raw-pixel OCR case while the JPEG case passes. Also on ef67cc3fa, one copy that drops the
PDF info dictionary, joins single-mode pages with a plain newline, refuses every encrypted PDF (with
a message that fails the password test too), no longer recognises pypandoc's "No pandoc was found"
and has a pyarrow that cannot build an array (patched in at import) fails the page-by-page,
single-mode, editing-lock, pandoc and every workbook test; a copy that opens a locked PDF with its
owner password fails the password test. A non-PDF fails its upload whether or not pypdf raises, so
that test only shows the refusal. An xlrd that lists only a workbook's first sheet and shifts every
date by a day (patched in at import) fails both legacy workbook tests, and one whose parse error is
replaced by a generic one fails the broken record test. A `detect_filetype` answering "unknown" for
every file (patched in at import) fails the legacy Word, legacy PowerPoint and Outlook tests, and a
loader that takes every .doc and .ppt for legacy Office fails the tests of modern files under those
names.
"""

from __future__ import annotations

import io
import shutil
import struct
import subprocess
import zipfile
import zlib
from pathlib import Path

import httpx
import pytest

from harness.actors import admin_of, create_user
from harness.listener import json_answer
from harness.missing_packages import without_packages_env
from harness.outlook_message import outlook_message

pytestmark = [pytest.mark.depcheck, pytest.mark.api, pytest.mark.requires_source]

RETRIEVAL_CONFIG = ("/api/v1/retrieval/config", "/api/v1/retrieval/config/update")
FIXTURES = Path(__file__).parent / "fixtures"

MARKER = "harbour lighthouse budget"
SENTENCE = f"The {MARKER} was approved."
OCR_WORD = "LIGHTHOUSE"

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
PPTX_TYPE = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _upload(client: httpx.Client, filename: str, content: bytes, content_type: str) -> str:
    """Upload and process a file before answering; returns its id."""
    uploaded = client.post(
        "/api/v1/files/",
        params={"process": "true", "process_in_background": "false"},
        files={"file": (filename, content, content_type)},
    )
    assert uploaded.status_code == 200, uploaded.text
    return uploaded.json()["id"]


def _read_back(client: httpx.Client, file_id: str) -> str:
    stored = client.get(f"/api/v1/files/{file_id}").json()
    assert stored["data"].get("status") == "completed", (
        f"{stored['filename']} was not read: {stored['data'].get('error')}"
    )
    read = client.get(f"/api/v1/files/{file_id}/data/content")
    assert read.status_code == 200, read.text
    return read.json()["content"]


def _upload_and_read(client: httpx.Client, filename: str, content: bytes, content_type: str) -> str:
    return _read_back(client, _upload(client, filename, content, content_type))


# ---------------------------------------------------------------- fixtures built in code


def _saved(document) -> bytes:
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _text_pdf(blank_first_page: bool = False) -> bytes:
    """One page carrying the sentence in Helvetica, the smallest PDF pypdf reads text from."""
    stream = f"BT /F1 12 Tf 72 720 Td ({SENTENCE}) Tj ET".encode()
    pages = b"[6 0 R 3 0 R] /Count 2" if blank_first_page else b"[3 0 R] /Count 1"
    return _assembled_pdf(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids %s >>" % pages,
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
            b"/Resources << /Font << /F1 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
            b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] >>",
        ]
    )


def _assembled_pdf(objects: list[bytes], info: bytes = b"") -> bytes:
    """A PDF of the numbered objects, the first the catalog, with its cross-reference table."""
    if info:
        objects = [*objects, info]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref_offset = len(pdf)
    pdf += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    pdf += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    info_entry = b" /Info %d 0 R" % len(objects) if info else b""
    pdf += b"trailer\n<< /Size %d /Root 1 0 R%s >>\n" % (len(objects) + 1, info_entry)
    pdf += b"startxref\n%d\n%%%%EOF\n" % xref_offset
    return bytes(pdf)


def _pdf(pages: list[str], info: bytes = b"", page_labels: bytes = b"") -> bytes:
    """One page per text in Helvetica, with an optional info dictionary and page label tree."""
    labels = b" /PageLabels %s" % page_labels if page_labels else b""
    page_numbers = [4 + 2 * index for index in range(len(pages))]
    kids = b" ".join(b"%d 0 R" % number for number in page_numbers)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R%s >>" % labels,
        b"<< /Type /Pages /Kids [%s] /Count %d >>" % (kids, len(pages)),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    for number, text in zip(page_numbers, pages):
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents %d 0 R "
            b"/Resources << /Font << /F1 3 0 R >> >> >>" % (number + 1)
        )
        objects.append(b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream))
    return _assembled_pdf(objects, info)


def _drawn(word: str):
    from PIL import Image, ImageDraw, ImageFont

    image = Image.new("L", (900, 200), 255)
    ImageDraw.Draw(image).text((30, 60), word, fill=0, font=ImageFont.load_default(size=64))
    return image


def _image_pdf(word: str) -> bytes:
    """A scan: the word drawn into an image, saved by Pillow as a PDF (a JPEG, no text layer)."""
    buffer = io.BytesIO()
    _drawn(word).save(buffer, format="PDF")
    return buffer.getvalue()


def _raw_image_pdf(word: str) -> bytes:
    """The same scan kept as compressed raw pixels, which Pillow cannot open without pypdf."""
    pixels = zlib.compress(_drawn(word).tobytes())
    stream = b"q 900 0 0 200 0 0 cm /Scan Do Q"
    return _assembled_pdf(
        [
            b"<< /Type /Catalog /Pages 2 0 R >>",
            b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 900 200] /Contents 4 0 R "
            b"/Resources << /XObject << /Scan 5 0 R >> >> >>",
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream),
            b"<< /Type /XObject /Subtype /Image /Width 900 /Height 200 /ColorSpace /DeviceGray "
            b"/BitsPerComponent 8 /Filter /FlateDecode /Length %d >>\nstream\n%s\nendstream"
            % (len(pixels), pixels),
        ]
    )


def _docx() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph(SENTENCE)
    return _saved(document)


def _pptx() -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Harbour review"
    slide.placeholders[1].text = SENTENCE
    return _saved(presentation)


def _xlsx() -> bytes:
    import openpyxl

    workbook = openpyxl.Workbook()
    workbook.active.append(["item", "note"])
    workbook.active.append(["lighthouse", MARKER])
    return _saved(workbook)


def _workbook() -> bytes:
    """Two sheets, the first with a whole and a fractional number."""
    import openpyxl

    workbook = openpyxl.Workbook()
    budget = workbook.active
    budget.title = "Budget"
    for row in (["item", "amount"], ["lighthouse", 1250], ["ferry", 80.5]):
        budget.append(row)
    crew = workbook.create_sheet("Crew")
    for row in (["name", "role"], ["Mara", "keeper"]):
        crew.append(row)
    return _saved(workbook)


def _xls() -> bytes:
    # xlwt, the only .xls writer, is no dependency; LibreOffice saved the rows `_xlsx` writes.
    return (FIXTURES / "budget.xls").read_bytes()


def _biff12_record(record_id: int, payload: bytes = b"") -> bytes:
    """One record of a binary workbook part: its id, a 7-bit length and the payload."""
    # an id keeps the high bit of each of its bytes, as the format numbers them
    head = bytes([record_id]) if record_id < 0x80 else struct.pack("<H", record_id)
    size, length = len(payload), bytearray()
    while True:
        length.append((size & 0x7F) | (0x80 if size > 0x7F else 0))
        size >>= 7
        if not size:
            break
    return head + bytes(length) + payload


def _wide_string(text: str) -> bytes:
    return struct.pack("<I", len(text)) + text.encode("utf-16-le")


def _xlsb(rows: list[list[str | float]]) -> bytes:
    """A binary workbook (.xlsb) of one sheet, which Excel writes and pyxlsb only reads."""
    record = _biff12_record
    strings = sorted({value for row in rows for value in row if isinstance(value, str)})
    cells = b""
    for row_number, row in enumerate(rows):
        cells += record(0x0000, struct.pack("<I", row_number))
        for column, value in enumerate(row):
            if isinstance(value, str):
                cells += record(0x0007, struct.pack("<III", column, 0, strings.index(value)))
            else:
                cells += record(0x0005, struct.pack("<IId", column, 0, value))
    extent = struct.pack("<IIII", 0, len(rows) - 1, 0, max(map(len, rows)) - 1)
    sheet = record(0x0181) + record(0x0194, extent) + record(0x0191) + cells + record(0x0192)
    listing = record(0x019C, struct.pack("<II", 0, 1) + _wide_string("rId1") + _wide_string("Log"))
    workbook = record(0x0183) + record(0x018F) + listing + record(0x0190) + record(0x0184)
    shared = record(0x019F, struct.pack("<II", len(strings), len(strings)))
    shared += b"".join(record(0x0013, b"\x00" + _wide_string(text)) for text in strings)
    content_types = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">'
        '<Default Extension="bin"'
        ' ContentType="application/vnd.ms-excel.sheet.binary.macroEnabled.main"/>'
        "</Types>"
    )
    relationships = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        '<Relationship Id="rId1" Target="worksheets/sheet1.bin" Type="http://schemas.openxmlformats.org'
        '/officeDocument/2006/relationships/worksheet"/></Relationships>'
    )
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("[Content_Types].xml", content_types)
        archive.writestr("xl/workbook.bin", workbook)
        archive.writestr("xl/_rels/workbook.bin.rels", relationships)
        archive.writestr("xl/worksheets/sheet1.bin", sheet + record(0x0182))
        archive.writestr("xl/sharedStrings.bin", shared + record(0x01A0))
    return buffer.getvalue()


def _ledger_xls() -> bytes:
    """Two sheets saved by LibreOffice as .xls: text, whole and fractional numbers, and dates."""
    return (FIXTURES / "ledger.xls").read_bytes()


LEDGER_ROWS = [
    ["lighthouse", "1250", "2024-03-15"],
    ["ferry", "80.5", "2024-11-02"],
]


def _xml() -> bytes:
    return f"<?xml version='1.0'?><notes><note>{SENTENCE}</note></notes>".encode()


def _html() -> bytes:
    page = f"<html><head><title>Harbour</title></head><body><p>{SENTENCE}</p></body></html>"
    return page.encode()


def _rst() -> bytes:
    return f"Harbour\n=======\n\n{SENTENCE}\n".encode()


def _container(mimetype: str, entries: dict[str, str]) -> bytes:
    """An EPUB or OpenDocument zip: the `mimetype` entry first and uncompressed."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("mimetype", mimetype, compress_type=zipfile.ZIP_STORED)
        for name, text in entries.items():
            archive.writestr(name, text)
    return buffer.getvalue()


def _epub() -> bytes:
    container = (
        '<?xml version="1.0"?>'
        '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/>'
        "</rootfiles></container>"
    )
    package = (
        '<?xml version="1.0"?>'
        '<package xmlns="http://www.idpf.org/2007/opf" version="2.0" unique-identifier="id">'
        '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>Harbour</dc:title>'
        '<dc:identifier id="id">harbour</dc:identifier><dc:language>en</dc:language></metadata>'
        '<manifest><item id="text" href="text.xhtml" media-type="application/xhtml+xml"/>'
        '</manifest><spine><itemref idref="text"/></spine></package>'
    )
    chapter = (
        '<html xmlns="http://www.w3.org/1999/xhtml"><head><title>Harbour</title></head>'
        f"<body><p>{SENTENCE}</p></body></html>"
    )
    return _container(
        "application/epub+zip",
        {"META-INF/container.xml": container, "content.opf": package, "text.xhtml": chapter},
    )


def _odt() -> bytes:
    manifest = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0"'
        ' manifest:version="1.2">'
        '<manifest:file-entry manifest:full-path="/"'
        ' manifest:media-type="application/vnd.oasis.opendocument.text"/>'
        '<manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/>'
        '<manifest:file-entry manifest:full-path="styles.xml" manifest:media-type="text/xml"/>'
        "</manifest:manifest>"
    )
    styles = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-styles xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' office:version="1.2"><office:styles/></office:document-styles>'
    )
    content = (
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"'
        ' xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2">'
        f"<office:body><office:text><text:p>{SENTENCE}</text:p></office:text></office:body>"
        "</office:document-content>"
    )
    return _container(
        "application/vnd.oasis.opendocument.text",
        {"META-INF/manifest.xml": manifest, "content.xml": content, "styles.xml": styles},
    )


def _require_pandoc() -> None:
    import pypandoc

    try:
        pypandoc.get_pandoc_version()
    except OSError:
        pytest.skip("no pandoc binary; unstructured converts .rst, .epub and .odt with it")


# ---------------------------------------------------------------- one upload per format


@pytest.mark.parametrize(
    ("filename", "build", "content_type"),
    [
        pytest.param("notes.pdf", _text_pdf, "application/pdf", id="pdf"),
        pytest.param("notes.docx", _docx, DOCX_TYPE, id="docx"),
        pytest.param("slides.pptx", _pptx, PPTX_TYPE, id="pptx"),
        pytest.param("budget.xlsx", _xlsx, XLSX_TYPE, id="xlsx"),
        pytest.param("budget.xls", _xls, "application/vnd.ms-excel", id="xls"),
        pytest.param(
            "budget.xlsb",
            lambda: _xlsb([["item", "note"], ["lighthouse", MARKER]]),
            "application/vnd.ms-excel",
            id="xlsb",
        ),
        pytest.param("notes.xml", _xml, "application/xml", id="xml"),
        pytest.param("notes.html", _html, "text/html", id="html"),
    ],
)
def test_an_uploaded_document_is_read_as_text(make_user, filename, build, content_type):
    with make_user().client() as client:
        content = _upload_and_read(client, filename, build(), content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"
    assert "<" not in content, f"{filename} kept its markup: {content!r}"


@pytest.mark.parametrize(
    ("filename", "build", "content_type"),
    [
        pytest.param("notes.rst", _rst, "text/x-rst", id="rst"),
        pytest.param("notes.epub", _epub, "application/epub+zip", id="epub"),
        pytest.param("notes.odt", _odt, "application/vnd.oasis.opendocument.text", id="odt"),
    ],
)
def test_a_document_pandoc_converts_is_read_as_text(make_user, filename, build, content_type):
    _require_pandoc()
    with make_user().client() as client:
        content = _upload_and_read(client, filename, build(), content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"


def test_every_sheet_of_a_workbook_is_read_in_order(make_user):
    with make_user().client() as client:
        content = _upload_and_read(client, "budget.xlsx", _workbook(), XLSX_TYPE)

    for row in ("lighthouse 1250", "ferry 80.5", "Mara keeper"):
        assert row in content, f"{row!r} is missing from {content!r}"
    assert content.index("lighthouse") < content.index("Mara")


def test_every_sheet_of_a_legacy_workbook_is_read_with_its_numbers_and_dates(make_user):
    with make_user().client() as client:
        content = _upload_and_read(client, "ledger.xls", _ledger_xls(), "application/vnd.ms-excel")

    for row in [*LEDGER_ROWS, ["Mara", "keeper"]]:
        assert " ".join(row) in content, f"{row} is missing from {content!r}"
    assert content.index("lighthouse") < content.index("Mara")


# ---------------------------------------------------------------- legacy Office and Outlook


@pytest.fixture(scope="module")
def legacy_office(tmp_path_factory):
    """(.doc, .ppt) saved by LibreOffice, which unstructured also converts them back with."""
    soffice = shutil.which("soffice")
    if soffice is None:
        pytest.skip("no soffice binary; unstructured reads .doc and .ppt through LibreOffice")
    directory = tmp_path_factory.mktemp("legacy-office")
    for name, build in (("notes.docx", _docx), ("slides.pptx", _pptx)):
        (directory / name).write_bytes(build())
        legacy = name[:-1]
        subprocess.run(
            [
                soffice,
                f"-env:UserInstallation=file://{directory}/profile",
                "--headless",
                "--convert-to",
                legacy.rsplit(".", 1)[1],
                "--outdir",
                str(directory),
                str(directory / name),
            ],
            check=True,
            capture_output=True,
            timeout=120,
        )
    return (directory / "notes.doc").read_bytes(), (directory / "slides.ppt").read_bytes()


@pytest.mark.parametrize(
    ("filename", "content_type", "which"),
    [
        pytest.param("notes.doc", "application/msword", 0, id="doc"),
        pytest.param("slides.ppt", "application/vnd.ms-powerpoint", 1, id="ppt"),
    ],
)
def test_a_legacy_office_document_is_read(make_user, legacy_office, filename, content_type, which):
    with make_user().client() as client:
        content = _upload_and_read(client, filename, legacy_office[which], content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"


@pytest.fixture(scope="module")
def without_soffice(instance_with, tmp_path_factory):
    """An instance on a host without LibreOffice: nothing on its PATH."""
    return instance_with({"PATH": str(tmp_path_factory.mktemp("empty-path"))})


@pytest.mark.slow
@pytest.mark.parametrize(
    ("filename", "build", "content_type"),
    [
        pytest.param("notes.doc", _docx, "application/msword", id="docx-as-doc"),
        pytest.param("slides.ppt", _pptx, "application/vnd.ms-powerpoint", id="pptx-as-ppt"),
    ],
)
def test_a_modern_office_file_under_a_legacy_name_is_read_by_its_content(
    without_soffice, filename, build, content_type
):
    # only a file taken for legacy Office needs LibreOffice
    with create_user(without_soffice).client() as client:
        content = _upload_and_read(client, filename, build(), content_type)

    assert MARKER in content, f"{filename} was read as {content!r}"


def test_an_outlook_message_is_read_as_its_body(make_user):
    message = outlook_message(
        subject="Harbour budget",
        body=f"{SENTENCE}\nRegards, Ingrid",
        sender_name="Ingrid Holm",
        sender_address="ingrid@harbour.example",
    )
    with make_user().client() as client:
        content = _upload_and_read(client, "budget.msg", message, "application/vnd.ms-outlook")

    assert content == f"{SENTENCE}\n\nRegards, Ingrid", content


# ---------------------------------------------------------------- without unstructured


@pytest.fixture(scope="module")
def without_unstructured(instance_with, tmp_path_factory):
    """An instance installed without the optional `unstructured` extra."""
    directory = tmp_path_factory.mktemp("without-unstructured")
    return instance_with(without_packages_env(directory, ["unstructured"]))


@pytest.mark.slow
def test_without_unstructured_every_sheet_is_read_by_pandas(without_unstructured):
    with create_user(without_unstructured).client() as client:
        content = _upload_and_read(client, "budget.xlsx", _workbook(), XLSX_TYPE)

    sheets = content.split("\n\n")
    assert [sheet.splitlines()[0] for sheet in sheets] == ["Sheet: Budget", "Sheet: Crew"]
    budget_rows = [line.split() for line in sheets[0].splitlines()[1:]]
    # no row index column: each row starts with its first cell
    assert budget_rows == [["item", "amount"], ["lighthouse", "1250.0"], ["ferry", "80.5"]]
    assert [line.split() for line in sheets[1].splitlines()[1:]] == [
        ["name", "role"],
        ["Mara", "keeper"],
    ]


@pytest.mark.slow
def test_without_unstructured_an_xls_is_read_by_pandas_on_xlrd(without_unstructured):
    with create_user(without_unstructured).client() as client:
        content = _upload_and_read(client, "budget.xls", _xls(), "application/vnd.ms-excel")

    assert content.startswith("Sheet: Budget\n"), content
    assert MARKER in content


@pytest.mark.slow
def test_without_unstructured_a_legacy_workbook_is_read_sheet_by_sheet(without_unstructured):
    with create_user(without_unstructured).client() as client:
        content = _upload_and_read(client, "ledger.xls", _ledger_xls(), "application/vnd.ms-excel")

    sheets = content.split("\n\n")
    assert [sheet.splitlines()[0] for sheet in sheets] == ["Sheet: Budget", "Sheet: Crew"]
    budget_rows = [line.split() for line in sheets[0].splitlines()[2:]]
    assert budget_rows == [
        ["lighthouse", "1250.0", "2024-03-15"],
        ["ferry", "80.5", "2024-11-02"],
    ]
    assert "Mara keeper" in sheets[1]


def _without_first_bof(workbook: bytes) -> bytes:
    """The workbook with its first BIFF8 beginning-of-file record blanked, its container intact."""
    start = workbook.index(b"\x09\x08\x10\x00\x00\x06")
    return workbook[:start] + b"\x00\x00" + workbook[start + 2 :]


@pytest.mark.slow
def test_without_unstructured_a_legacy_workbook_xlrd_cannot_parse_is_refused(without_unstructured):
    with create_user(without_unstructured).client() as client:
        uploaded = _upload(
            client, "ledger.xls", _without_first_bof(_ledger_xls()), "application/vnd.ms-excel"
        )
        stored = client.get(f"/api/v1/files/{uploaded}").json()

    assert stored["data"].get("status") == "failed", stored["data"]
    assert "Expected BOF record" in stored["data"].get("error", ""), stored["data"]


@pytest.mark.slow
def test_without_unstructured_slides_are_read_by_python_pptx(without_unstructured):
    with create_user(without_unstructured).client() as client:
        content = _upload_and_read(client, "slides.pptx", _pptx(), PPTX_TYPE)

    assert content == f"Slide 1:\nHarbour review\n{SENTENCE}", content


def test_a_binary_workbook_keeps_its_numbers(make_user):
    rows = [["harbour", "high water"], ["Northgate", 4.25], ["Southpier", 17.5]]
    with make_user().client() as client:
        content = _upload_and_read(client, "tides.xlsb", _xlsb(rows), "application/vnd.ms-excel")

    assert "Northgate 4.25" in content and "Southpier 17.5" in content, content


# ---------------------------------------------------------------- text encodings


@pytest.mark.parametrize(
    ("codec", "text"),
    [
        pytest.param("gb18030", "港口灯塔的预算已经批准。简体中文编码测试。", id="gb18030"),
        pytest.param("big5", "港口燈塔的預算已經批准。繁體中文編碼測試。", id="big5"),
        pytest.param(
            "euc-kr", "항구 등대 예산이 승인되었습니다. 한국어 인코딩 감지 테스트.", id="euc-kr"
        ),
        pytest.param(
            "shift_jis",
            "港の灯台の予算が承認されました。日本語エンコーディング検出テスト。",
            id="shift-jis",
        ),
    ],
)
def test_a_cjk_text_file_is_decoded(make_user, codec, text):
    document = text * 8
    with make_user().client() as client:
        content = _upload_and_read(client, "notes.txt", document.encode(codec), "text/plain")

    assert content == document, f"{codec} was decoded as {content[:40]!r}"


def test_a_cyrillic_windows_1251_file_is_decoded_by_chardets_guess(make_user):
    # no CJK codec reads it, so the loader takes chardet's own answer before latin-1
    document = "Смотритель маяка записывает приливы и погоду в судовой журнал. " * 4
    with make_user().client() as client:
        content = _upload_and_read(client, "notes.txt", document.encode("cp1251"), "text/plain")

    assert content == document, f"cp1251 was decoded as {content[:40]!r}"


def test_a_pdf_with_a_blank_page_is_read(make_user):
    with make_user().client() as client:
        content = _upload_and_read(
            client, "notes.pdf", _text_pdf(blank_first_page=True), "application/pdf"
        )

    assert MARKER in content, f"the PDF was read as {content!r}"


# ---------------------------------------------------------------- Word documents


WORD_PARAGRAPHS = [
    "Grüße aus dem Hafen, café au lait am Kai.",
    "港口灯塔的预算已经批准。",
    "Маяк работает всю ночь.",
    "The harbour lighthouse budget was approved.",
]


def test_a_word_document_keeps_every_paragraph_in_order(make_user):
    import docx

    document = docx.Document()
    for paragraph in WORD_PARAGRAPHS:
        document.add_paragraph(paragraph)
    with make_user().client() as client:
        content = _upload_and_read(client, "letters.docx", _saved(document), DOCX_TYPE)

    positions = [content.find(paragraph) for paragraph in WORD_PARAGRAPHS]
    assert -1 not in positions, f"a paragraph was lost or garbled: {content!r}"
    assert positions == sorted(positions), f"the paragraphs came back out of order: {content!r}"


def test_an_empty_word_document_is_reported_as_empty(make_user):
    import docx

    with make_user().client() as client:
        file_id = _upload(client, "empty.docx", _saved(docx.Document()), DOCX_TYPE)
        stored = client.get(f"/api/v1/files/{file_id}").json()

    assert stored["data"]["status"] == "failed", stored["data"]
    assert "empty" in stored["data"]["error"], stored["data"]


# ---------------------------------------------------------------- text repair


def test_mojibake_is_repaired_and_a_literal_entity_is_kept(make_user):
    # UTF-8 text once decoded as Windows-1252, so its apostrophe became "â€™".
    garbled = "The harbour lighthouse budget wasnâ€™t cut &amp; the keeper stays."
    with make_user().client() as client:
        content = _upload_and_read(client, "notes.txt", garbled.encode(), "text/plain")

    assert "budget wasn't cut" in content, f"the mojibake survived: {content!r}"
    assert "&amp;" in content, f"the literal entity was unescaped: {content!r}"


def _straightened(text: str) -> str:
    return text.translate(str.maketrans({"“": '"', "”": '"', "‘": "'", "’": "'"}))


def test_smart_quotes_read_as_latin_1_are_repaired(make_user):
    intended = "He said “hello” and ‘goodbye’ to the harbour master."
    garbled = intended.encode("utf-8").decode("latin-1")
    with make_user().client() as client:
        content = _upload_and_read(client, "quotes.txt", garbled.encode(), "text/plain")

    # ftfy also straightens curly quotes by default, which is fine
    assert _straightened(content) == _straightened(intended), f"not repaired: {content!r}"


def test_control_characters_and_terminal_colours_are_removed(make_user):
    with make_user().client() as client:
        content = _upload_and_read(
            client,
            "log.txt",
            "The harbour\x00 lighthouse \x1b[31mbudget\x1b[0m was\x07 approved.".encode(),
            "text/plain",
        )

    assert content == SENTENCE, f"control characters survived: {content!r}"


# ---------------------------------------------------------------- OCR of PDF images


@pytest.fixture
def retrieval_settings(preserve, admin):
    """`update(**settings)` changes the document settings for this test."""
    preserve(RETRIEVAL_CONFIG)
    client = admin.client()

    def update(**settings) -> None:
        updated = client.post(RETRIEVAL_CONFIG[1], json=settings)
        assert updated.status_code == 200, updated.text

    yield update
    client.close()


DEAD_PROXY = "http://127.0.0.1:9"
# every download refused, as on a machine without internet
OFFLINE_OCR_ENV = {
    "PDF_EXTRACT_IMAGES": "true",
    **{name: DEAD_PROXY for name in ("HTTP_PROXY", "HTTPS_PROXY", "http_proxy", "https_proxy")},
    **{name: "127.0.0.1,localhost" for name in ("NO_PROXY", "no_proxy")},
}


@pytest.fixture
def offline_ocr(instance_with):
    """An instance reading PDF images with rapidocr and no way out to the internet."""
    return admin_of(instance_with(OFFLINE_OCR_ENV))


@pytest.mark.slow
@pytest.mark.parametrize("build", [_image_pdf, _raw_image_pdf], ids=["jpeg-image", "raw-pixels"])
def test_the_text_in_a_pdf_image_is_read_without_downloading_anything(offline_ocr, build):
    scan = build(OCR_WORD)
    with offline_ocr.client() as client:
        content = _upload_and_read(client, "scan.pdf", scan, "application/pdf")

    assert OCR_WORD in content, f"OCR read {content!r}"


@pytest.mark.slow
def test_a_pdf_image_without_text_adds_nothing(offline_ocr):
    with offline_ocr.client() as client:
        file_id = _upload(client, "blank-scan.pdf", _image_pdf(""), "application/pdf")
        stored = client.get(f"/api/v1/files/{file_id}").json()
        read = client.get(f"/api/v1/files/{file_id}/data/content")

    assert read.json()["content"] == ""
    assert stored["data"]["status"] == "failed"
    assert "content provided is empty" in stored["data"]["error"]


def test_without_image_extraction_a_scan_yields_no_text(retrieval_settings, make_user):
    retrieval_settings(PDF_EXTRACT_IMAGES=False)
    with make_user().client() as client:
        file_id = _upload(client, "scan.pdf", _image_pdf(OCR_WORD), "application/pdf")
        read = client.get(f"/api/v1/files/{file_id}/data/content")

    assert OCR_WORD not in read.json()["content"]


# ---------------------------------------------------------------- Azure Document Intelligence

AZURE_MODEL = "prebuilt-layout"
AZURE_KEY = "document-intelligence-key"
AZURE_ANALYZE_PATH = f"/documentintelligence/documentModels/{AZURE_MODEL}:analyze"
AZURE_RESULT_PATH = f"/documentintelligence/documentModels/{AZURE_MODEL}/analyzeResults/one"
AZURE_MARKDOWN = f"# Harbour review\n\n{SENTENCE}\n"


def _serve_document_intelligence(listener) -> None:
    """Azure's analyze call: accepted with an operation to poll, which has already succeeded."""
    operation = f"{listener.base_url}{AZURE_RESULT_PATH}?api-version=2024-11-30"
    accepted = {"Operation-Location": operation, "Retry-After": "0"}
    listener.route("POST", AZURE_ANALYZE_PATH, (202, accepted, b""))
    result = {
        "status": "succeeded",
        "createdDateTime": "2026-01-01T00:00:00Z",
        "lastUpdatedDateTime": "2026-01-01T00:00:01Z",
        "analyzeResult": {
            "apiVersion": "2024-11-30",
            "modelId": AZURE_MODEL,
            "content": AZURE_MARKDOWN,
            "contentFormat": "markdown",
            "pages": [],
        },
    }
    listener.route("GET", AZURE_RESULT_PATH, json_answer(result))


def test_document_intelligence_reads_a_pdf_as_markdown(retrieval_settings, make_user, listener):
    _serve_document_intelligence(listener)
    retrieval_settings(
        CONTENT_EXTRACTION_ENGINE="document_intelligence",
        DOCUMENT_INTELLIGENCE_ENDPOINT=listener.base_url,
        DOCUMENT_INTELLIGENCE_KEY=AZURE_KEY,
        DOCUMENT_INTELLIGENCE_MODEL=AZURE_MODEL,
    )
    pdf = _text_pdf()

    with make_user().client() as client:
        content = _upload_and_read(client, "review.pdf", pdf, "application/pdf")

    assert content.strip() == AZURE_MARKDOWN.strip()
    [analyze] = listener.requests_to(AZURE_ANALYZE_PATH)
    assert analyze.headers["Ocp-Apim-Subscription-Key"] == AZURE_KEY
    assert "outputContentFormat=markdown" in analyze.path
    assert analyze.body == pdf
    assert listener.requests_to(AZURE_RESULT_PATH), "the analysis result was never polled"


# ---------------------------------------------------------------- password-protected workbooks


def _encrypted_xlsx() -> bytes:
    from msoffcrypto.format.ooxml import OOXMLFile

    encrypted = io.BytesIO()
    OOXMLFile(io.BytesIO(_xlsx())).encrypt("harbour-password", encrypted)
    return encrypted.getvalue()


def _processing_error(client: httpx.Client, filename: str, content: bytes) -> str:
    """Why the upload could not be read, from the file's stored status."""
    stored = client.get(f"/api/v1/files/{_upload(client, filename, content, XLSX_TYPE)}").json()
    assert stored["data"].get("status") != "completed", f"{filename} was read: {stored}"
    return str(stored["data"].get("error"))


def test_a_password_protected_workbook_is_refused_as_such(make_user):
    with make_user().client() as client:
        error = _processing_error(client, "locked.xlsx", _encrypted_xlsx())

    assert "password protected" in error, error


def test_a_file_that_only_claims_to_be_a_workbook_is_refused(make_user):
    with make_user().client() as client:
        error = _processing_error(client, "fake.xlsx", b"not a workbook, only its name")

    assert "Not a valid XLSX file" in error, error


# ---------------------------------------------------------------- PDFs, page by page

HARBOUR_PAGES = ["The harbour chart shows the channel.", f"The {MARKER} was approved."]
SURVEY_INFO = (
    b"<< /Title (  Harbour survey  ) /Author (The keeper) "
    b"/CreationDate (D:20240102030405+00'00') >>"
)
# the first page is numbered in roman numerals, the rest from 1
ROMAN_THEN_ARABIC = b"<< /Nums [0 << /S /r >> 1 << /S /D >>] >>"
PDF_PASSWORD = "harbour-password"


def _chunks(client: httpx.Client, file_id: str) -> list[tuple[str, dict]]:
    """Every stored chunk of the file with its metadata, in page order."""
    found = client.post(
        "/api/v1/retrieval/query/doc",
        json={"collection_name": f"file-{file_id}", "query": "harbour", "k": 20},
    )
    assert found.status_code == 200, found.text
    chunks = zip(found.json()["documents"][0], found.json()["metadatas"][0])
    return sorted(chunks, key=lambda chunk: chunk[1].get("page", 0))


def _locked_pdf(user_password: str) -> bytes:
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter(clone_from=PdfReader(io.BytesIO(_pdf(HARBOUR_PAGES))))
    writer.encrypt(user_password=user_password, owner_password="harbour-owner")
    locked = io.BytesIO()
    writer.write(locked)
    return locked.getvalue()


def test_a_pdf_is_read_page_by_page_with_its_labels_and_metadata(make_user):
    pdf = _pdf(HARBOUR_PAGES, info=SURVEY_INFO, page_labels=ROMAN_THEN_ARABIC)
    with make_user().client() as client:
        file_id = _upload(client, "survey.pdf", pdf, "application/pdf")
        content = _read_back(client, file_id)
        chunks = _chunks(client, file_id)

    assert content == " ".join(HARBOUR_PAGES)
    assert [text for text, _ in chunks] == HARBOUR_PAGES
    assert [(meta["page"], meta["page_label"]) for _, meta in chunks] == [(0, "i"), (1, "1")]
    first = chunks[0][1]
    assert first["title"] == "Harbour survey", first
    assert first["author"] == "The keeper"
    assert first["creationdate"] == "2024-01-02T03:04:05+00:00"
    assert first["total_pages"] == 2


def test_in_single_mode_a_pdf_is_one_document_split_at_page_breaks(retrieval_settings, make_user):
    retrieval_settings(PDF_LOADER_MODE="single")
    with make_user().client() as client:
        content = _upload_and_read(client, "survey.pdf", _pdf(HARBOUR_PAGES), "application/pdf")

    assert content == "\n\f".join(HARBOUR_PAGES)


def test_a_pdf_locked_only_against_editing_is_read(make_user):
    with make_user().client() as client:
        content = _upload_and_read(client, "survey.pdf", _locked_pdf(""), "application/pdf")

    assert MARKER in content, content


def test_a_pdf_that_needs_a_password_to_open_is_refused(make_user):
    with make_user().client() as client:
        file_id = _upload(client, "survey.pdf", _locked_pdf(PDF_PASSWORD), "application/pdf")
        stored = client.get(f"/api/v1/files/{file_id}").json()
        read = client.get(f"/api/v1/files/{file_id}/data/content")

    assert stored["data"].get("status") == "failed", stored["data"]
    assert "decrypt" in str(stored["data"].get("error")).lower(), stored["data"]
    assert MARKER not in read.text


def test_a_file_that_only_claims_to_be_a_pdf_is_refused(make_user):
    with make_user().client() as client:
        file_id = _upload(client, "survey.pdf", b"not a pdf, only its name", "application/pdf")
        stored = client.get(f"/api/v1/files/{file_id}").json()

    assert stored["data"].get("status") == "failed", stored["data"]
    assert stored["data"].get("error"), stored["data"]


# ---------------------------------------------------------------- without pandoc

PANDOC_MISSING = (
    "Pandoc is not installed on the server. Please contact your administrator for assistance."
)


def test_without_pandoc_a_document_it_converts_names_the_missing_program(make_user):
    import pypandoc

    try:
        pypandoc.get_pandoc_version()
    except OSError:
        pass
    else:
        pytest.skip("pandoc is installed; this is the answer an install without it gives")
    with make_user().client() as client:
        file_id = _upload(client, "notes.rst", _rst(), "text/x-rst")
        stored = client.get(f"/api/v1/files/{file_id}").json()

    assert stored["data"].get("status") == "failed", stored["data"]
    assert stored["data"]["error"] == PANDOC_MISSING, stored["data"]
